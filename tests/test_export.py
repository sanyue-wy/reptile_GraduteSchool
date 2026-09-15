# -*- coding: utf-8 -*-
"""
导出管道单元测试
"""

import json
from pathlib import Path

import pytest

from pipelines.export import (
    export_merged,
    export_summary,
    export_failures,
    load_failures,
    add_failure,
    update_failure_status,
    get_active_failures,
    clear_failure,
    ensure_output_dir,
)


class TestExportMerged:
    """合并数据 JSONL 导出"""

    def test_export_merged_jsonl(self, tmp_dir):
        """输出 JSONL 每行可解析，字段完整"""
        records = [
            {
                "schema_version": 1,
                "name": "张三",
                "university": "测试大学",
                "college": "机械工程学院",
                "category": "mechanical",
                "title": "教授",
                "advisor_level": "博导",
                "email": "zhangsan@test.edu.cn",
                "profile_url": "http://test.edu.cn/1",
                "research_areas": ["智能制造"],
                "enrollment": None,
                "match_status": "partial_faculty",
                "raw_ref": {"faculty": "", "notice": ""},
                "source_type": {"faculty": "官网师资页", "notice": ""},
            },
            {
                "schema_version": 1,
                "name": "李四",
                "university": "测试大学",
                "college": "机械工程学院",
                "category": "mechanical",
                "title": "副教授",
                "advisor_level": "硕导",
                "email": "lisi@test.edu.cn",
                "profile_url": "http://test.edu.cn/2",
                "research_areas": ["机器人学"],
                "enrollment": None,
                "match_status": "partial_faculty",
                "raw_ref": {"faculty": "", "notice": ""},
                "source_type": {"faculty": "官网师资页", "notice": ""},
            },
        ]

        result = export_merged(records, output_dir=tmp_dir)
        assert result["files_written"] == 1
        assert result["total_records"] == 2

        # 验证文件名和内容
        jsonl_files = list(tmp_dir.glob("*.jsonl"))
        assert len(jsonl_files) == 1
        lines = jsonl_files[0].read_text(encoding="utf-8").strip().split("\n")
        assert len(lines) == 2
        for line in lines:
            rec = json.loads(line)
            assert "name" in rec
            assert "university" in rec
            assert "match_status" in rec

    def test_export_merged_empty(self, tmp_dir):
        """空记录列表不写入文件"""
        result = export_merged([], output_dir=tmp_dir)
        assert result["files_written"] == 0
        assert result["total_records"] == 0

    def test_export_merged_different_schools(self, tmp_dir):
        """不同学校/学院分组写入不同文件"""
        records = [
            {"name": "张三", "university": "大学A", "college": "学院A1", "match_status": "merged",
             "research_areas": [], "title": "", "advisor_level": "", "email": "",
             "profile_url": "", "enrollment": None, "raw_ref": {}, "source_type": {}},
            {"name": "李四", "university": "大学B", "college": "学院B1", "match_status": "merged",
             "research_areas": [], "title": "", "advisor_level": "", "email": "",
             "profile_url": "", "enrollment": None, "raw_ref": {}, "source_type": {}},
        ]
        result = export_merged(records, output_dir=tmp_dir)
        assert result["files_written"] == 2
        assert result["total_records"] == 2


class TestExportSummary:
    """Excel 汇总表导出"""

    def test_export_summary_excel(self, tmp_dir):
        """summary.xlsx 可打开，含正确列头和数据行"""
        import openpyxl
        records = [
            {
                "schema_version": 1,
                "name": "张三",
                "university": "测试大学",
                "college": "机械工程学院",
                "title": "教授",
                "advisor_level": "博导",
                "email": "zhangsan@test.edu.cn",
                "profile_url": "http://test.edu.cn/1",
                "research_areas": ["智能制造", "机器人学"],
                "enrollment": {
                    "in_roster": True,
                    "degree_types": ["硕士"],
                    "directions": [{"code": "085501", "name": "机械工程"}],
                    "planned_count": 3,
                    "exam_subjects": ["政治"],
                    "source_url": "",
                    "notice_year": 2026,
                },
                "match_status": "merged",
                "raw_ref": {"faculty": "", "notice": ""},
                "source_type": {"faculty": "官网师资页", "notice": "研招网"},
            },
        ]

        output_path = tmp_dir / "summary.xlsx"
        count = export_summary(records, output_path=output_path)
        assert count == 1

        # 验证 Excel 文件
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
        headers = [ws.cell(row=1, column=col).value for col in range(1, 9)]
        assert "学校" in headers
        assert "学院" in headers
        assert "姓名" in headers
        assert "职称" in headers
        assert "招生状态" in headers
        assert "招生方向" in headers
        assert "研究方向" in headers
        assert "来源链接" in headers

        # 验证数据行
        row2 = [ws.cell(row=2, column=col).value for col in range(1, 9)]
        assert "张三" in row2
        assert "教授" in row2


class TestExportFailures:
    """失败记录导出"""

    def test_export_failures_json(self, tmp_dir):
        """failures.json 可解析，结构符合规范"""
        failures = [
            add_failure(
                failure_id="fail_001",
                school="测试大学",
                college="机械工程学院",
                source="Source A",
                error_type="http_error",
                error_message="403 Forbidden",
                url="http://test.edu.cn/faculty",
                retry_count=1,
            ),
        ]
        output_path = tmp_dir / "failures.json"
        count = export_failures(failures, output_path=output_path)
        assert count == 1

        data = json.loads(output_path.read_text(encoding="utf-8"))
        assert data["total"] == 1
        assert "summary" in data
        assert data["summary"]["http_error"] == 1
        assert len(data["items"]) == 1
        assert data["items"][0]["id"] == "fail_001"

    def test_load_failures_empty(self, tmp_dir):
        """无 failures.json 时返回空列表"""
        output_path = tmp_dir / "failures.json"
        result = load_failures(input_path=output_path)
        assert result == []

    def test_load_failures_existing(self, tmp_dir):
        """读取现有 failures.json"""
        failures = [
            add_failure("fail_001", "大学A", "学院A", "Source A", "timeout", "timeout error", "http://test.com"),
        ]
        output_path = tmp_dir / "failures.json"
        export_failures(failures, output_path=output_path)

        result = load_failures(input_path=output_path)
        assert len(result) == 1
        assert result[0]["id"] == "fail_001"

    def test_update_failure_status(self, tmp_dir):
        """更新失败记录状态"""
        failures = [
            add_failure("fail_001", "大学A", "学院A", "Source A", "timeout", "error", "http://test.com"),
        ]
        output_path = tmp_dir / "failures.json"
        export_failures(failures, output_path=output_path)

        result = update_failure_status("fail_001", "resolved", output_path=output_path)
        assert result is True

        loaded = load_failures(input_path=output_path)
        assert loaded[0]["status"] == "resolved"

    def test_update_failure_status_not_found(self, tmp_dir):
        """更新不存在的失败记录返回 False"""
        output_path = tmp_dir / "failures.json"
        assert update_failure_status("nonexistent", "resolved", output_path=output_path) is False

    def test_get_active_failures(self, tmp_dir):
        """获取 active 状态的失败记录"""
        failures = [
            add_failure("fail_001", "大学A", "学院A", "Source A", "timeout", "error", "http://test.com", status="active"),
            add_failure("fail_002", "大学B", "学院B", "Source A", "http_error", "error", "http://test.com", status="resolved"),
            add_failure("fail_003", "大学C", "学院C", "Source A", "parse_error", "error", "http://test.com", status="ignored"),
        ]
        output_path = tmp_dir / "failures.json"
        export_failures(failures, output_path=output_path)

        active = get_active_failures(output_path=output_path)
        assert len(active) == 1
        assert active[0]["id"] == "fail_001"

    def test_clear_failure(self, tmp_dir):
        """删除一条失败记录"""
        failures = [
            add_failure("fail_001", "大学A", "学院A", "Source A", "timeout", "error", "http://test.com"),
            add_failure("fail_002", "大学B", "学院B", "Source A", "timeout", "error", "http://test.com"),
        ]
        output_path = tmp_dir / "failures.json"
        export_failures(failures, output_path=output_path)

        result = clear_failure("fail_001", output_path=output_path)
        assert result is True

        loaded = load_failures(input_path=output_path)
        assert len(loaded) == 1
        assert loaded[0]["id"] == "fail_002"

    def test_clear_failure_not_found(self, tmp_dir):
        """删除不存在的记录返回 False"""
        output_path = tmp_dir / "failures.json"
        assert clear_failure("nonexistent", output_path=output_path) is False

    def test_add_failure(self):
        """创建失败记录结构"""
        record = add_failure(
            failure_id="fail_test",
            school="测试大学",
            college="机械学院",
            source="Source A",
            error_type="http_403",
            error_message="Forbidden",
            url="http://test.com",
        )
        assert record["id"] == "fail_test"
        assert record["school"] == "测试大学"
        assert record["college"] == "机械学院"
        assert record["source"] == "Source A"
        assert record["error_type"] == "http_403"
        assert record["error_message"] == "Forbidden"
        assert record["url"] == "http://test.com"
        assert record["status"] == "active"
        assert "occurred_at" in record
