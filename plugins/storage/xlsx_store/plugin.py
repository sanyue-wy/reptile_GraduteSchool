"""XLSX Storage Plugin.

Wraps Excel export functionality with V3.0 plugin interface.
"""

import logging
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from contracts.result import StoreRequest, StoreReceipt, ErrorDTO
from plugins.base import BasePlugin, PluginContext

logger = logging.getLogger(__name__)


class XlsxStorePlugin(BasePlugin[StoreRequest, StoreReceipt]):
    """Storage plugin for XLSX output."""

    name = "xlsx_store"
    version = "1.0.0"
    plugin_type = "storage"
    input_schema = "StoreRequest.v1"
    output_schema = "StoreReceipt.v1"

    HEADERS = ["学校", "学院", "姓名", "职称", "招生状态", "招生方向", "研究方向", "来源链接"]
    COLUMN_WIDTHS = [18, 18, 12, 10, 10, 40, 40, 50]

    def __init__(self):
        super().__init__()
        self._config = {}

    def setup(self, context: PluginContext) -> None:
        self._context = context
        self._config = context.config_snapshot.get("plugins", {}).get("xlsx_store", {})

    def execute(self, request: StoreRequest, context: PluginContext) -> StoreReceipt:
        try:
            from infra.storage import check_and_reserve, commit as idem_commit
            from converters.storage_converter import (
                records_from_data_refs, records_to_legacy_format,
            )

            # Idempotency check
            base_dir = Path(self._config.get("output_dir", "data"))
            if context.storage and hasattr(context.storage, "workspace"):
                base_dir = context.storage.workspace.base_dir
            cached = check_and_reserve(base_dir, request.run_id, request.idempotency_key)
            if cached is not None:
                return StoreReceipt(
                    target_id=cached.get("target_id", request.target_id),
                    written=0,
                    skipped=cached.get("records_written", 0),
                    failed=0,
                    records_written=0,
                    output_ref=cached.get("output_ref", ""),
                )

            output_path = self._resolve_output_path(request, context)
            raw_records = records_from_data_refs(request.data_refs)

            if not raw_records:
                self._create_empty_workbook(output_path)
                receipt = StoreReceipt(
                    target_id=request.target_id,
                    written=0,
                    skipped=0,
                    failed=0,
                    records_written=0,
                    output_ref=str(output_path),
                )
            else:
                records = records_to_legacy_format(raw_records)
                written = self._write_xlsx(records, output_path)
                receipt = StoreReceipt(
                    target_id=request.target_id,
                    written=written,
                    skipped=0,
                    failed=0,
                    records_written=written,
                    output_ref=str(output_path),
                )

            idem_commit(base_dir, request.run_id, request.idempotency_key, {
                "target_id": receipt.target_id,
                "records_written": receipt.records_written,
                "output_ref": receipt.output_ref,
            })

            return receipt

        except Exception as e:
            logger.exception("XLSX storage failed for target %s", request.target_id)
            return StoreReceipt(
                target_id=request.target_id,
                written=0,
                skipped=0,
                failed=1,
                records_written=0,
                output_ref="",
                error=ErrorDTO(
                    code="STORAGE_WRITE_FAILED",
                    message=str(e),
                    stage="store",
                    task_id=request.state_snapshot.get("task_id", ""),
                ),
            )

    def _resolve_output_path(self, request: StoreRequest, context: PluginContext) -> Path:
        output_dir = self._config.get("output_dir", "data/output")
        filename = self._config.get("filename", "summary.xlsx")
        base = Path(output_dir)
        base.mkdir(parents=True, exist_ok=True)
        return base / filename

    def _create_empty_workbook(self, output_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = self._config.get("sheet_name", "导师汇总")
        ws.append(self.HEADERS)
        self._style_header(ws)
        self._style_columns(ws)
        from infra.storage.atomic_io import atomic_write_bytes
        bio = BytesIO()
        wb.save(bio)
        atomic_write_bytes(output_path, bio.getvalue())

    def _write_xlsx(self, records: list[dict], output_path: Path) -> int:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = self._config.get("sheet_name", "导师汇总")
        ws.append(self.HEADERS)
        self._style_header(ws)

        for row_idx, rec in enumerate(records, start=2):
            row_data = self._record_to_row(rec)
            ws.append(row_data)
            for col_idx in range(1, len(self.HEADERS) + 1):
                ws.cell(row=row_idx, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")

        self._style_columns(ws)
        ws.freeze_panes = "A2"

        from infra.storage.atomic_io import atomic_write_bytes
        bio = BytesIO()
        wb.save(bio)
        atomic_write_bytes(output_path, bio.getvalue())

        logger.info("Wrote %d records to %s", len(records), output_path)
        return len(records)

    def _record_to_row(self, rec: dict) -> list[str]:
        return [
            rec.get("university", ""),
            rec.get("college", ""),
            rec.get("name", ""),
            rec.get("title", ""),
            rec.get("advisor_status", ""),
            rec.get("research_areas", ""),
            rec.get("research_areas", ""),
            rec.get("source_url", ""),
        ]

    def _style_header(self, ws) -> None:
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        for col_idx, _header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _style_columns(self, ws) -> None:
        for i, width in enumerate(self.COLUMN_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

    def close(self) -> None:
        pass